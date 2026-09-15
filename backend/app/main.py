from fastapi import FastAPI, File, UploadFile, Depends
from fastapi.responses import JSONResponse, FileResponse, HTMLResponse
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
import shutil
import os
import json
from app.extraction import extract_invoice_data
from app.export import export_to_excel
from app.tax_engine import process_invoice_tax, classify_transaction
from app.models import create_tables, get_db, Invoice

app = FastAPI(
    title="Taxora API",
    description="AI-powered bookkeeping and tax-preparation platform for Pakistani businesses",
    version="2.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

create_tables()

FRONTEND = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "frontend", "index.html")
if os.path.exists(FRONTEND):
    with open(FRONTEND, "r", encoding="utf-8") as f:
        HTML = f.read()
else:
    HTML = "<h1>Frontend not found</h1>"

@app.get("/")
def root():
    return HTMLResponse(content=HTML)

@app.get("/health")
def health():
    return {"status": "healthy", "version": "2.0.0"}

@app.post("/upload-and-extract")
async def upload_and_extract(file: UploadFile = File(...), db: Session = Depends(get_db)):
    allowed_types = ["image/jpeg", "image/png"]
    if file.content_type not in allowed_types:
        return JSONResponse(status_code=400, content={"error": "Only JPG and PNG files are supported"})
    
    file_path = os.path.join(UPLOAD_DIR, file.filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    
    extracted = extract_invoice_data(file_path, file.content_type)
    tax_result = process_invoice_tax(extracted)
    
    invoice = Invoice(
        filename=file.filename,
        supplier_name=extracted.get("supplier_name", ""),
        invoice_number=extracted.get("invoice_number", ""),
        invoice_date=extracted.get("invoice_date", ""),
        due_date=extracted.get("due_date", ""),
        description=extracted.get("description", ""),
        subtotal=extracted.get("subtotal", ""),
        sales_tax=extracted.get("sales_tax", ""),
        withholding_tax=extracted.get("withholding_tax", ""),
        total_amount=extracted.get("total_amount", ""),
        currency=extracted.get("currency", ""),
        payment_method=extracted.get("payment_method", ""),
        transaction_type=tax_result.get("transaction_type", "other"),
        tax_rate=tax_result.get("rate_applied", ""),
        tax_amount=tax_result.get("withholding_tax", 0.0),
        fbr_section=tax_result.get("section", ""),
        status="extracted"
    )
    db.add(invoice)
    db.commit()
    db.refresh(invoice)
    
    return {
        "message": "Invoice uploaded and extracted successfully",
        "invoice_id": invoice.id,
        "extracted_data": extracted,
        "tax_calculation": tax_result
    }

@app.get("/invoices")
def get_invoices(db: Session = Depends(get_db)):
    invoices = db.query(Invoice).order_by(Invoice.created_at.desc()).all()
    return [
        {
            "id": inv.id,
            "filename": inv.filename,
            "supplier_name": inv.supplier_name,
            "invoice_number": inv.invoice_number,
            "invoice_date": inv.invoice_date,
            "total_amount": inv.total_amount,
            "currency": inv.currency,
            "transaction_type": inv.transaction_type,
            "tax_rate": inv.tax_rate,
            "tax_amount": inv.tax_amount,
            "fbr_section": inv.fbr_section,
            "status": inv.status,
            "created_at": inv.created_at.isoformat()
        }
        for inv in invoices
    ]

@app.get("/invoices/{invoice_id}")
def get_invoice(invoice_id: int, db: Session = Depends(get_db)):
    inv = db.query(Invoice).filter(Invoice.id == invoice_id).first()
    if not inv:
        return JSONResponse(status_code=404, content={"error": "Invoice not found"})
    return {
        "id": inv.id,
        "filename": inv.filename,
        "supplier_name": inv.supplier_name,
        "invoice_number": inv.invoice_number,
        "invoice_date": inv.invoice_date,
        "due_date": inv.due_date,
        "description": inv.description,
        "subtotal": inv.subtotal,
        "sales_tax": inv.sales_tax,
        "withholding_tax": inv.withholding_tax,
        "total_amount": inv.total_amount,
        "currency": inv.currency,
        "payment_method": inv.payment_method,
        "transaction_type": inv.transaction_type,
        "tax_rate": inv.tax_rate,
        "tax_amount": inv.tax_amount,
        "fbr_section": inv.fbr_section,
        "status": inv.status,
        "created_at": inv.created_at.isoformat()
    }

@app.put("/invoices/{invoice_id}/approve")
def approve_invoice(invoice_id: int, db: Session = Depends(get_db)):
    inv = db.query(Invoice).filter(Invoice.id == invoice_id).first()
    if not inv:
        return JSONResponse(status_code=404, content={"error": "Invoice not found"})
    inv.status = "approved"
    db.commit()
    return {"message": "Invoice approved", "id": invoice_id}

@app.delete("/invoices/{invoice_id}")
def delete_invoice(invoice_id: int, db: Session = Depends(get_db)):
    inv = db.query(Invoice).filter(Invoice.id == invoice_id).first()
    if not inv:
        return JSONResponse(status_code=404, content={"error": "Invoice not found"})
    db.delete(inv)
    db.commit()
    return {"message": "Invoice deleted", "id": invoice_id}

@app.post("/generate-fbr-report")
async def generate_fbr_report(invoice_ids: list[int], db: Session = Depends(get_db)):
    invoices = db.query(Invoice).filter(Invoice.id.in_(invoice_ids)).all()
    if not invoices:
        return JSONResponse(status_code=404, content={"error": "No invoices found"})
    
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FBR Withholding Tax Statement"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    title_font = Font(bold=True, size=14, color="1F4E79")
    center = Alignment(horizontal="center", vertical="center")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    
    # Title
    ws.merge_cells("A1:J1")
    ws["A1"] = "WITHHOLDING TAX STATEMENT"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 35
    
    ws.merge_cells("A2:J2")
    ws["A2"] = "Federal Board of Revenue (FBR) - Pakistan | Tax Year 2027"
    ws["A2"].font = Font(bold=True, size=11, color="1F4E79")
    ws["A2"].alignment = center
    ws.row_dimensions[2].height = 25
    
    ws.merge_cells("A3:J3")
    ws["A3"] = "Generated by Taxora AI | www.taxora.ai"
    ws["A3"].font = Font(italic=True, size=9, color="666666")
    ws["A3"].alignment = center
    
    # Headers
    headers = [
        "S.No", "Supplier Name", "Invoice No.", "Invoice Date",
        "Description", "Gross Amount", "FBR Section",
        "Transaction Type", "WHT Rate", "WHT Amount"
    ]
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[5].height = 20
    
    # Data rows
    total_gross = 0
    total_wht = 0
    
    for row_num, (idx, inv) in enumerate(enumerate(invoices, start=1), start=6):
        try:
            amount = float(inv.total_amount or 0)
        except:
            amount = 0.0
        
        total_gross += amount
        total_wht += inv.tax_amount or 0
        
        row_data = [
            idx,
            inv.supplier_name or "-",
            inv.invoice_number or "-",
            inv.invoice_date or "-",
            inv.description or "-",
            amount,
            inv.fbr_section or "-",
            inv.transaction_type.replace("_", " ").title() if inv.transaction_type else "-",
            inv.tax_rate or "-",
            inv.tax_amount or 0
        ]
        
        fill = PatternFill("solid", fgColor="F5F9FF") if row_num % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.border = thin
            cell.alignment = Alignment(horizontal="center" if col in [1, 7, 8, 9] else "left")
            cell.fill = fill
    
    # Totals row
    total_row = len(invoices) + 6
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.merge_cells(f"A{total_row}:E{total_row}")
    ws.cell(row=total_row, column=6, value=total_gross).font = Font(bold=True)
    ws.cell(row=total_row, column=10, value=total_wht).font = Font(bold=True)
    
    for col in range(1, 11):
        ws.cell(row=total_row, column=col).fill = PatternFill("solid", fgColor="E8F0FE")
        ws.cell(row=total_row, column=col).border = thin
    
    # Column widths
    widths = [6, 25, 18, 14, 30, 15, 14, 20, 10, 14]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64+col)].width = width
    
    # Save
    export_dir = "exports"
    os.makedirs(export_dir, exist_ok=True)
    export_path = os.path.join(export_dir, "FBR_Withholding_Tax_Statement.xlsx")
    wb.save(export_path)
    
    return FileResponse(
        path=export_path,
        filename="FBR_Withholding_Tax_Statement.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.post("/extract")
async def extract_document(file: UploadFile = File(...)):
    allowed_types = ["image/jpeg", "image/png"]
    if file.content_type not in allowed_types:
        return JSONResponse(status_code=400, content={"error": "Only JPG and PNG files are supported"})
    file_path = os.path.join(UPLOAD_DIR, file.filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    extracted = extract_invoice_data(file_path, file.content_type)
    return {"message": "Extraction complete", "filename": file.filename, "extracted_data": extracted}