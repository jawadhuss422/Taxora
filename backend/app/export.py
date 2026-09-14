import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os

def export_to_excel(extracted_data: dict, filename: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice Data"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    center = Alignment(horizontal="center")

    # Title row
    ws.merge_cells("A1:C1")
    ws["A1"] = "TAXORA - Extracted Invoice Data"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = center

    # Column headers
    headers = ["Field", "Extracted Value"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # Field labels
    field_labels = {
        "supplier_name": "Supplier Name",
        "invoice_number": "Invoice Number",
        "invoice_date": "Invoice Date",
        "due_date": "Due Date",
        "description": "Description",
        "subtotal": "Subtotal",
        "sales_tax": "Sales Tax",
        "withholding_tax": "Withholding Tax",
        "total_amount": "Total Amount",
        "currency": "Currency",
        "payment_method": "Payment Method"
    }

    # Data rows
    for row, (key, label) in enumerate(field_labels.items(), start=4):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=extracted_data.get(key, ""))

    # Column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 35

    # Save
    export_dir = "exports"
    os.makedirs(export_dir, exist_ok=True)
    export_path = os.path.join(export_dir, f"{filename}_export.xlsx")
    wb.save(export_path)
    return export_path